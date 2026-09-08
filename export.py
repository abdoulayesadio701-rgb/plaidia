"""
export.py — Génère un document Word (.docx) ou PDF à partir d'une analyse,
pour que l'avocat puisse l'imprimer, l'annoter ou l'intégrer à son dossier.
"""

from datetime import datetime
from pathlib import Path
import paths

EXPORTS_DIR = paths.base_dir() / "exports"

# Page de garde : reproduction statique (PNG) du motif d'arche gothique de
# frontend/src/components/GothicMotif.tsx -- ce composant est du React/SVG,
# inutilisable tel quel dans un document Word/PDF généré côté Python, d'où
# cet asset généré une fois (voir assets/page_de_garde_motif.png) plutôt
# qu'un rendu à la volée. Absent = page de garde sans image, jamais une
# erreur (voir _ajouter_page_de_garde_word / _page_de_garde_pdf_flowables).
PAGE_DE_GARDE_MOTIF = paths.base_dir() / "assets" / "page_de_garde_motif.png"

RISK_ORDER = {"Élevé": 0, "Moyen": 1, "Faible": 2}


def _nom_fichier(dossier_nom: str, extension: str) -> str:
    safe_nom = "".join(c if c.isalnum() or c in " -_" else "_" for c in dossier_nom).strip()
    date_str = datetime.now().strftime("%Y-%m-%d_%Hh%M")
    return f"{safe_nom}_{date_str}.{extension}"


def _ajouter_page_de_garde_word(doc, titre: str, sous_titre: str = "") -> None:
    """Insère une page de garde (motif, titre, sous-titre, date) avant le
    contenu du document, puis un saut de page. Partagée par toutes les
    fonctions d'export Word ci-dessous."""
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    for _ in range(3):
        doc.add_paragraph()

    if PAGE_DE_GARDE_MOTIF.exists():
        p_image = doc.add_paragraph()
        p_image.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_image.add_run().add_picture(str(PAGE_DE_GARDE_MOTIF), width=Inches(2.1))

    p_titre = doc.add_paragraph()
    p_titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_titre = p_titre.add_run(titre)
    run_titre.bold = True
    run_titre.font.size = Pt(20)

    if sous_titre:
        p_sous = doc.add_paragraph()
        p_sous.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_sous = p_sous.add_run(sous_titre)
        run_sous.italic = True
        run_sous.font.size = Pt(12)

    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_date.add_run(f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}").italic = True
    p_date.runs[0].font.size = Pt(9)

    p_marque = doc.add_paragraph()
    p_marque.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_marque.add_run("Plaid'IA — assistant de préparation de plaidoirie").font.size = Pt(9)

    doc.add_page_break()


def _page_de_garde_pdf_flowables(titre: str, sous_titre: str = "") -> list:
    """Équivalent PDF de _ajouter_page_de_garde_word : une liste de
    flowables reportlab à préfixer à la 'story' de chaque export PDF."""
    from reportlab.platypus import Image, Paragraph, PageBreak, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib.colors import HexColor
    from reportlab.lib.enums import TA_CENTER

    styles = getSampleStyleSheet()
    titre_style = ParagraphStyle("PageGardeTitre", parent=styles["Title"], alignment=TA_CENTER, fontSize=20, spaceAfter=6)
    sous_style = ParagraphStyle("PageGardeSousTitre", parent=styles["Normal"], alignment=TA_CENTER, fontName="Helvetica-Oblique", fontSize=12)
    date_style = ParagraphStyle("PageGardeDate", parent=styles["Normal"], alignment=TA_CENTER, fontName="Helvetica-Oblique", fontSize=9, textColor=HexColor("#666666"))

    flowables = [Spacer(1, 3.5 * cm)]
    if PAGE_DE_GARDE_MOTIF.exists():
        img = Image(str(PAGE_DE_GARDE_MOTIF), width=3.5 * cm, height=5.25 * cm)
        img.hAlign = "CENTER"
        flowables += [img, Spacer(1, 1 * cm)]
    flowables.append(Paragraph(titre, titre_style))
    if sous_titre:
        flowables.append(Paragraph(sous_titre, sous_style))
    flowables += [
        Spacer(1, 0.6 * cm),
        Paragraph(f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}", date_style),
        Paragraph("Plaid'IA — assistant de préparation de plaidoirie", date_style),
        PageBreak(),
    ]
    return flowables


def exporter_word(dossier: dict, result: dict) -> str:
    """Génère un .docx et retourne le chemin du fichier créé."""
    dossier = dict(dossier)
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        raise ImportError("python-docx n'est pas installé. Lancez : pip install python-docx")

    EXPORTS_DIR.mkdir(exist_ok=True)

    doc = Document()
    _ajouter_page_de_garde_word(doc, f"Analyse — {dossier['nom']}", dossier.get("domaine") or "")

    title = doc.add_heading(f"Analyse — {dossier['nom']}", level=1)

    meta = doc.add_paragraph()
    meta.add_run(f"Domaine : {dossier.get('domaine') or 'non précisé'}\n").italic = True
    meta.add_run(f"Généré le : {datetime.now().strftime('%d/%m/%Y à %H:%M')}").italic = True

    doc.add_paragraph()

    args_sorted = sorted(
        result.get("arguments", []), key=lambda a: RISK_ORDER.get(a.get("risque", "Moyen"), 1)
    )

    for i, arg in enumerate(args_sorted, 1):
        h = doc.add_heading(f"{i}. {arg.get('resume', '')}", level=2)

        risk_p = doc.add_paragraph()
        risk_run = risk_p.add_run(f"Niveau de risque : {arg.get('risque', '?')}")
        risk_run.bold = True
        if arg.get("risque") == "Élevé":
            risk_run.font.color.rgb = RGBColor(0xC0, 0x30, 0x30)
        elif arg.get("risque") == "Moyen":
            risk_run.font.color.rgb = RGBColor(0xB0, 0x80, 0x00)
        else:
            risk_run.font.color.rgb = RGBColor(0x30, 0x80, 0x40)

        doc.add_paragraph(f"Fondement : {arg.get('fondement', '')}")
        raisonnement = arg.get("raisonnement") or {}
        if raisonnement.get("probleme_de_droit"):
            doc.add_paragraph(f"Problème de droit : {raisonnement.get('probleme_de_droit', '')}")
        if raisonnement.get("regle_applicable"):
            doc.add_paragraph(f"Règle applicable : {raisonnement.get('regle_applicable', '')}")
        if raisonnement.get("application_aux_faits"):
            doc.add_paragraph(f"Application aux faits : {raisonnement.get('application_aux_faits', '')}")
        doc.add_paragraph(f"Justification : {arg.get('justification_risque', '')}")

        if arg.get("refutations"):
            doc.add_paragraph("Pistes de réfutation :", style="Intense Quote")
            for r in arg["refutations"]:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(f"[{r.get('angle', '')}] ").bold = True
                p.add_run(r.get("piste", ""))

        doc.add_paragraph()

    if result.get("points_attention"):
        doc.add_heading("Points d'attention", level=2)
        for point in result["points_attention"]:
            doc.add_paragraph(point, style="List Bullet")

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer_run = footer.add_run(
        "Document généré par un outil d'assistance IA. Toute référence marquée "
        "\"À VÉRIFIER\" doit être contrôlée avant utilisation en plaidoirie."
    )
    footer_run.italic = True
    footer_run.font.size = Pt(9)

    filename = _nom_fichier(dossier["nom"], "docx")
    path = EXPORTS_DIR / filename
    doc.save(str(path))
    return str(path)


def exporter_pdf(dossier: dict, result: dict) -> str:
    """Génère un .pdf et retourne le chemin du fichier créé."""
    dossier = dict(dossier)
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import cm
        from reportlab.lib.colors import HexColor
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem
        )
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    except ImportError:
        raise ImportError("reportlab n'est pas installé. Lancez : pip install reportlab")

    EXPORTS_DIR.mkdir(exist_ok=True)

    filename = _nom_fichier(dossier["nom"], "pdf")
    path = EXPORTS_DIR / filename

    doc = SimpleDocTemplate(str(path), pagesize=A4,
                             leftMargin=2*cm, rightMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitrePerso", parent=styles["Heading1"], fontSize=16, spaceAfter=12)
    h2_style = ParagraphStyle("H2Perso", parent=styles["Heading2"], fontSize=12, spaceBefore=14, spaceAfter=6)
    normal = styles["Normal"]
    italic = ParagraphStyle("ItalicPerso", parent=normal, fontName="Helvetica-Oblique", fontSize=9, textColor=HexColor("#555555"))

    risk_colors = {"Élevé": "#C03030", "Moyen": "#B08000", "Faible": "#308040"}

    story = _page_de_garde_pdf_flowables(f"Analyse — {dossier['nom']}", dossier.get("domaine") or "")
    story.append(Paragraph(f"Analyse — {dossier['nom']}", title_style))
    story.append(Paragraph(f"Domaine : {dossier.get('domaine') or 'non précisé'}", italic))
    story.append(Paragraph(f"Généré le : {datetime.now().strftime('%d/%m/%Y à %H:%M')}", italic))
    story.append(Spacer(1, 0.5*cm))

    args_sorted = sorted(
        result.get("arguments", []), key=lambda a: RISK_ORDER.get(a.get("risque", "Moyen"), 1)
    )

    for i, arg in enumerate(args_sorted, 1):
        story.append(Paragraph(f"{i}. {arg.get('resume', '')}", h2_style))
        risque = arg.get("risque", "?")
        color = risk_colors.get(risque, "#333333")
        story.append(Paragraph(f'<font color="{color}"><b>Niveau de risque : {risque}</b></font>', normal))
        story.append(Paragraph(f"<b>Fondement :</b> {arg.get('fondement', '')}", normal))
        raisonnement = arg.get("raisonnement") or {}
        if raisonnement.get("probleme_de_droit"):
            story.append(Paragraph(f"<b>Problème de droit :</b> {raisonnement.get('probleme_de_droit', '')}", normal))
        if raisonnement.get("regle_applicable"):
            story.append(Paragraph(f"<b>Règle applicable :</b> {raisonnement.get('regle_applicable', '')}", normal))
        if raisonnement.get("application_aux_faits"):
            story.append(Paragraph(f"<b>Application aux faits :</b> {raisonnement.get('application_aux_faits', '')}", normal))
        story.append(Paragraph(f"<b>Justification :</b> {arg.get('justification_risque', '')}", normal))

        if arg.get("refutations"):
            items = []
            for r in arg["refutations"]:
                items.append(ListItem(Paragraph(f"<b>[{r.get('angle', '')}]</b> {r.get('piste', '')}", normal)))
            story.append(Spacer(1, 0.2*cm))
            story.append(ListFlowable(items, bulletType="bullet"))

        story.append(Spacer(1, 0.4*cm))

    if result.get("points_attention"):
        story.append(Paragraph("Points d'attention", h2_style))
        items = [ListItem(Paragraph(p, normal)) for p in result["points_attention"]]
        story.append(ListFlowable(items, bulletType="bullet"))

    story.append(Spacer(1, 0.6*cm))
    story.append(Paragraph(
        "Document généré par un outil d'assistance IA. Toute référence marquée "
        "« À VÉRIFIER » doit être contrôlée avant utilisation en plaidoirie.",
        italic
    ))

    doc.build(story)
    return str(path)


def exporter_dossier_complet_word(dossier: dict, analyse: dict | None, plan: dict | None, simulateur: dict | None) -> str:
    """Génère un .docx unique combinant analyse, plan de plaidoirie et
    simulateur d'objections pour un dossier — le rapport complet à emporter
    à l'audience. Chaque section est optionnelle."""
    dossier = dict(dossier)
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
    except ImportError:
        raise ImportError("python-docx n'est pas installé. Lancez : pip install python-docx")

    EXPORTS_DIR.mkdir(exist_ok=True)
    doc = Document()
    _ajouter_page_de_garde_word(doc, f"Dossier complet — {dossier['nom']}", dossier.get("domaine") or "")

    doc.add_heading(f"Dossier complet — {dossier['nom']}", level=1)
    meta = doc.add_paragraph()
    meta.add_run(f"Domaine : {dossier.get('domaine') or 'non précisé'}\n").italic = True
    meta.add_run(f"Généré le : {datetime.now().strftime('%d/%m/%Y à %H:%M')}").italic = True
    doc.add_paragraph()

    if plan:
        doc.add_heading("1. Plan de plaidoirie", level=1)
        doc.add_paragraph("Accroche", style="Intense Quote")
        doc.add_paragraph(plan.get("accroche", ""))
        for i, point in enumerate(plan.get("plan", []), 1):
            doc.add_heading(f"{i}. {point.get('point', '')} ({point.get('duree_minutes', '?')} min)", level=2)
            doc.add_paragraph(f"Argument clé : {point.get('argument_cle', '')}")
            doc.add_paragraph(f"Notes : {point.get('notes', '')}")
        doc.add_paragraph("Conclusion", style="Intense Quote")
        doc.add_paragraph(plan.get("conclusion", ""))
        if plan.get("points_attention"):
            doc.add_heading("Points d'attention (plan)", level=2)
            for p in plan["points_attention"]:
                doc.add_paragraph(p, style="List Bullet")
        doc.add_paragraph()

    if analyse:
        doc.add_heading("2. Analyse des arguments adverses", level=1)
        args_sorted = sorted(analyse.get("arguments", []), key=lambda a: RISK_ORDER.get(a.get("risque", "Moyen"), 1))
        for i, arg in enumerate(args_sorted, 1):
            h = doc.add_heading(f"{i}. {arg.get('resume', '')}", level=2)
            risk_p = doc.add_paragraph()
            risk_run = risk_p.add_run(f"Niveau de risque : {arg.get('risque', '?')}")
            risk_run.bold = True
            doc.add_paragraph(f"Fondement : {arg.get('fondement', '')}")
            raisonnement = arg.get("raisonnement") or {}
            if raisonnement.get("probleme_de_droit"):
                doc.add_paragraph(f"Problème de droit : {raisonnement.get('probleme_de_droit', '')}")
            if raisonnement.get("regle_applicable"):
                doc.add_paragraph(f"Règle applicable : {raisonnement.get('regle_applicable', '')}")
            if raisonnement.get("application_aux_faits"):
                doc.add_paragraph(f"Application aux faits : {raisonnement.get('application_aux_faits', '')}")
            if arg.get("justification_risque"):
                doc.add_paragraph(f"Justification : {arg.get('justification_risque', '')}")
            if arg.get("refutations"):
                doc.add_paragraph("Pistes de réfutation :", style="Intense Quote")
                for r in arg["refutations"]:
                    p = doc.add_paragraph(style="List Bullet")
                    p.add_run(f"[{r.get('angle', '')}] ").bold = True
                    p.add_run(r.get("piste", ""))
        if analyse.get("points_attention"):
            doc.add_heading("Points d'attention (analyse)", level=2)
            for p in analyse["points_attention"]:
                doc.add_paragraph(p, style="List Bullet")
        doc.add_paragraph()

    if simulateur:
        doc.add_heading("3. Questions et objections probables", level=1)
        for i, obj in enumerate(simulateur.get("objections", []), 1):
            doc.add_heading(f"{i}. [{obj.get('origine', '?')}] {obj.get('question', '')}", level=2)
            doc.add_paragraph(f"Piège : {obj.get('piege', '')}")
            doc.add_paragraph(f"Piste de réponse : {obj.get('piste_reponse', '')}")
        if simulateur.get("point_le_plus_faible"):
            doc.add_heading("Point le plus faible du dossier", level=2)
            doc.add_paragraph(simulateur["point_le_plus_faible"])
        doc.add_paragraph()

    footer = doc.add_paragraph()
    footer_run = footer.add_run(
        "Document généré par un outil d'assistance IA. Toute référence marquée "
        "\"À VÉRIFIER\" doit être contrôlée avant utilisation en plaidoirie."
    )
    footer_run.italic = True
    footer_run.font.size = Pt(9)

    filename = _nom_fichier(f"{dossier['nom']}_complet", "docx")
    path = EXPORTS_DIR / filename
    doc.save(str(path))
    return str(path)


def exporter_note_client_word(dossier: dict, texte_note: str) -> str:
    """Génère un .docx contenant la note explicative destinée au client,
    en langage simple — prête à être envoyée telle quelle."""
    dossier = dict(dossier)
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        raise ImportError("python-docx n'est pas installé. Lancez : pip install python-docx")

    EXPORTS_DIR.mkdir(exist_ok=True)
    doc = Document()
    _ajouter_page_de_garde_word(doc, f"Point sur votre dossier — {dossier['nom']}")

    doc.add_heading(f"Point sur votre dossier — {dossier['nom']}", level=1)
    meta = doc.add_paragraph()
    meta.add_run(f"Préparé le {datetime.now().strftime('%d/%m/%Y')}").italic = True
    doc.add_paragraph()

    for paragraphe in texte_note.split("\n\n"):
        if paragraphe.strip():
            doc.add_paragraph(paragraphe.strip())

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer_run = footer.add_run(
        "Ce document est une synthèse simplifiée destinée à votre information. "
        "N'hésitez pas à contacter votre avocat pour toute question."
    )
    footer_run.italic = True
    footer_run.font.size = Pt(9)

    filename = _nom_fichier(f"{dossier['nom']}_note_client", "docx")
    path = EXPORTS_DIR / filename
    doc.save(str(path))
    return str(path)


def exporter_faits_bruts_word(dossier: dict) -> str:
    """Exporte les faits bruts d'un dossier (tels qu'accumulés via
    'Préparer mon dossier' ou saisis à la création), sans analyse ni
    reformulation — pour classer avec le dossier physique du client."""
    dossier = dict(dossier)
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        raise ImportError("python-docx n'est pas installé. Lancez : pip install python-docx")

    EXPORTS_DIR.mkdir(exist_ok=True)
    doc = Document()
    _ajouter_page_de_garde_word(doc, f"Faits du dossier — {dossier['nom']}", dossier.get("domaine") or "")

    doc.add_heading(f"Faits du dossier — {dossier['nom']}", level=1)
    meta = doc.add_paragraph()
    meta.add_run(f"Domaine : {dossier.get('domaine') or 'non précisé'}\n").italic = True
    if dossier.get("numero_dossier"):
        meta.add_run(f"Numéro de référence : {dossier['numero_dossier']}\n").italic = True
    if dossier.get("parties"):
        meta.add_run(f"Parties : {dossier['parties']}\n").italic = True
    meta.add_run(f"Exporté le : {datetime.now().strftime('%d/%m/%Y à %H:%M')}").italic = True
    doc.add_paragraph()

    faits = dossier.get("faits") or ""
    if not faits.strip():
        doc.add_paragraph("Aucun fait enregistré pour ce dossier.")
    else:
        for ligne in faits.split("\n"):
            if ligne.strip():
                doc.add_paragraph(ligne.strip())
            else:
                doc.add_paragraph()

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer_run = footer.add_run(
        "Export brut des faits saisis dans l'outil — aucune analyse ni reformulation IA."
    )
    footer_run.italic = True
    footer_run.font.size = Pt(9)

    filename = _nom_fichier(f"{dossier['nom']}_faits_bruts", "docx")
    path = EXPORTS_DIR / filename
    doc.save(str(path))
    return str(path)


def exporter_texte_libre_word(titre: str, texte: str, note_bas_page: str = "") -> str:
    """Export générique d'un texte libre en Word, sans nécessiter de
    dossier en base — utilisé notamment pour les procès-verbaux, qui
    peuvent être rédigés avant même la création d'une affaire dans l'outil."""
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        raise ImportError("python-docx n'est pas installé. Lancez : pip install python-docx")

    EXPORTS_DIR.mkdir(exist_ok=True)
    doc = Document()
    _ajouter_page_de_garde_word(doc, titre)

    doc.add_heading(titre, level=1)
    meta = doc.add_paragraph()
    meta.add_run(f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}").italic = True
    doc.add_paragraph()

    for ligne in texte.split("\n"):
        if ligne.strip():
            doc.add_paragraph(ligne.strip())
        else:
            doc.add_paragraph()

    if note_bas_page:
        doc.add_paragraph()
        footer = doc.add_paragraph()
        footer_run = footer.add_run(note_bas_page)
        footer_run.italic = True
        footer_run.font.size = Pt(9)

    filename = _nom_fichier(titre, "docx")
    path = EXPORTS_DIR / filename
    doc.save(str(path))
    return str(path)


def exporter_csv(titre: str, en_tetes: list[str], lignes: list[list[str]]) -> str:
    """Export générique d'un tableau en CSV — pour les résultats
    naturellement tabulaires (ex. chronologie) où un tableur est plus
    approprié qu'un document Word (voir AUDIT_IMPORT_EXPORT.md §6).
    UTF-8 avec BOM (utf-8-sig) pour qu'Excel affiche correctement les
    accents à l'ouverture directe du fichier, sans réglage manuel."""
    import csv

    EXPORTS_DIR.mkdir(exist_ok=True)
    filename = _nom_fichier(titre, "csv")
    path = EXPORTS_DIR / filename
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(en_tetes)
        writer.writerows(lignes)
    return str(path)

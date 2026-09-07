#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
tests/test_extract.py — Tests unitaires pour le module extract.py
"""

import pytest
import tempfile
from pathlib import Path


class TestExtractText:
    """Tests pour l'extraction de texte."""
    
    def test_extract_txt(self):
        """Test l'extraction de texte depuis un fichier TXT."""
        import sys
        sys.path.insert(0, str(Path(__file__).parent.parent))
        
        from extract import extract_text
        
        # Créer un fichier TXT temporaire
        with tempfile.NamedTemporaryFile(mode='w', suffix='.txt', delete=False, encoding='utf-8') as f:
            f.write("Ceci est un texte de test.\nAvec plusieurs lignes.\n")
            temp_path = f.name
        
        try:
            text = extract_text(temp_path)
            assert "texte de test" in text
            assert "plusieurs lignes" in text
        finally:
            Path(temp_path).unlink()
    
    def test_extract_txt_unicode(self):
        """Test l'extraction de texte avec caractères spéciaux."""
        import sys
        sys.path.insert(0, str(Path(__file__).parent.parent))
        
        from extract import extract_text
        
        # Créer un fichier TXT avec caractères français
        with tempfile.NamedTemporaryFile(mode='w', suffix='.txt', delete=False, encoding='utf-8') as f:
            f.write("Analyse juridique : prud'hommes, contrat, rémunération.\n")
            temp_path = f.name
        
        try:
            text = extract_text(temp_path)
            assert "prud'hommes" in text
            assert "rémunération" in text
        finally:
            Path(temp_path).unlink()


class TestExtractPdf:
    """Tests pour l'extraction PDF et la détection de document numérisé (§8)."""

    def _generer_pdf(self, chemin, texte_par_page):
        """Génère un vrai PDF via reportlab (déjà une dépendance du projet,
        utilisée par export.py) -- un `texte_par_page` vide produit une
        page blanche, sans couche de texte, comme un scan sans OCR."""
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas

        c = canvas.Canvas(str(chemin), pagesize=A4)
        for texte in texte_par_page:
            if texte:
                y = 750
                for ligne in texte.split("\n"):
                    c.drawString(100, y, ligne)
                    y -= 20
            c.showPage()
        c.save()

    def test_extract_pdf_normal(self, tmp_path):
        import sys
        sys.path.insert(0, str(Path(__file__).parent.parent))
        from extract import extract_text

        chemin = tmp_path / "normal.pdf"
        self._generer_pdf(chemin, ["Ceci est un vrai document juridique avec du texte exploitable.\nIl contient largement plus de vingt caracteres par page."])

        texte = extract_text(str(chemin))
        assert "document juridique" in texte

    def test_extract_pdf_scanne_leve_document_numerise_error(self, tmp_path):
        """Un PDF sans texte exploitable (page blanche, comme un scan sans
        OCR) doit lever DocumentNumeriseError, pas renvoyer une chaîne vide
        silencieusement -- voir ARCHITECTURE_CHAT_CONTEXTUEL.md §2.6/§8."""
        import sys
        sys.path.insert(0, str(Path(__file__).parent.parent))
        from extract import DocumentNumeriseError, extract_text

        chemin = tmp_path / "scan.pdf"
        self._generer_pdf(chemin, [""])

        with pytest.raises(DocumentNumeriseError, match="numérisé"):
            extract_text(str(chemin))

    def test_extract_pdf_plusieurs_pages_avec_peu_de_texte_leve_aussi(self, tmp_path):
        """Le seuil est proportionnel au nombre de pages : quelques mots
        perdus dans un document de 5 pages doivent aussi être détectés,
        pas seulement une page totalement vide."""
        import sys
        sys.path.insert(0, str(Path(__file__).parent.parent))
        from extract import DocumentNumeriseError, extract_text

        chemin = tmp_path / "scan_5_pages.pdf"
        self._generer_pdf(chemin, ["Titre"] + [""] * 4)

        with pytest.raises(DocumentNumeriseError):
            extract_text(str(chemin))


class TestExtractDocx:
    def test_extract_docx_reel(self, tmp_path):
        import sys
        sys.path.insert(0, str(Path(__file__).parent.parent))
        import docx
        from extract import extract_text

        chemin = tmp_path / "conclusions.docx"
        document = docx.Document()
        document.add_paragraph("Conclusions récapitulatives du défendeur.")
        document.add_paragraph("Il est demandé au tribunal de débouter le demandeur.")
        document.save(str(chemin))

        texte = extract_text(str(chemin))
        assert "Conclusions récapitulatives" in texte
        assert "débouter le demandeur" in texte


class TestExtractXlsx:
    def test_extract_xlsx_reel(self, tmp_path):
        import sys
        sys.path.insert(0, str(Path(__file__).parent.parent))
        import openpyxl
        from extract import extract_text

        chemin = tmp_path / "prejudice.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Calcul du préjudice"
        ws.append(["Poste", "Montant"])
        ws.append(["Salaire impayé", 4500])
        wb.save(str(chemin))

        texte = extract_text(str(chemin))
        assert "Calcul du préjudice" in texte
        assert "Salaire impayé" in texte
        assert "4500" in texte


class TestCleanText:
    """Tests pour le nettoyage de texte."""
    
    def test_clean_text_removes_blank_lines(self):
        """Vérifie que _clean_text supprime les lignes vides."""
        import sys
        sys.path.insert(0, str(Path(__file__).parent.parent))
        
        from extract import _clean_text
        
        text = "Ligne 1\n\n\nLigne 2\n\n\nLigne 3"
        cleaned = _clean_text(text)
        
        # Pas plus de 1 ligne vide consécutive
        assert "\n\n\n" not in cleaned
    
    def test_clean_text_preserves_content(self):
        """Vérifie que _clean_text préserve le contenu."""
        import sys
        sys.path.insert(0, str(Path(__file__).parent.parent))
        
        from extract import _clean_text
        
        text = "Contenu important à préserver"
        cleaned = _clean_text(text)
        
        assert "Contenu important" in cleaned


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
